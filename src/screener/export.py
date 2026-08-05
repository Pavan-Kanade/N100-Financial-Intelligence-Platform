"""
Screener Excel Exporter Module for Nifty 100 Financial Intelligence Platform.
Generates output/screener_output.xlsx with 6 sheets and openpyxl cell color-coding.
"""

import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pandas as pd
from typing import Dict

from src.screener.engine import ScreenerEngine


class ScreenerExporter:
    def __init__(self, output_dir: str = "output"):
        self.output_dir = output_dir
        os.makedirs(self.output_dir, exist_ok=True)
        self.export_path = os.path.join(self.output_dir, "screener_output.xlsx")

    def export(self, preset_results: Dict[str, pd.DataFrame]):
        """
        Exports preset results to screener_output.xlsx with openpyxl styling.
        """
        wb = openpyxl.Workbook()

        # Styles
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")

        green_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
        red_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

        border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        cols_to_export = [
            "company_id", "company_name", "broad_sector", "composite_quality_score",
            "return_on_equity_pct", "debt_to_equity", "free_cash_flow_cr",
            "revenue_cagr_5yr", "pat_cagr_5yr", "operating_profit_margin_pct",
            "pe_ratio", "pb_ratio", "dividend_yield_pct", "interest_coverage",
            "market_cap_crore", "net_profit", "sales", "eps_cagr_5yr",
            "asset_turnover", "book_value_per_share"
        ]

        # Remove default sheet
        default_sheet = wb.active

        for idx, (preset_name, df) in enumerate(preset_results.items()):
            ws = wb.create_sheet(title=preset_name)

            # Filter columns that exist
            valid_cols = [c for c in cols_to_export if c in df.columns]
            df_export = df[valid_cols].copy()

            # Write headers
            for col_num, col_name in enumerate(valid_cols, 1):
                cell = ws.cell(row=1, column=col_num)
                cell.value = col_name
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Write data rows
            for row_idx, row_data in enumerate(df_export.to_dict("records"), 2):
                for col_idx, col_name in enumerate(valid_cols, 1):
                    val = row_data.get(col_name)
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.value = val
                    cell.border = border

                    # Cell alignment
                    if isinstance(val, (int, float)):
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="center")

                    # Cell color-coding logic
                    if isinstance(val, (int, float)):
                        if col_name in ["return_on_equity_pct", "free_cash_flow_cr", "revenue_cagr_5yr", "pat_cagr_5yr"]:
                            cell.fill = green_fill if val > 0 else red_fill
                        elif col_name == "debt_to_equity":
                            cell.fill = green_fill if val <= 1.0 else (red_fill if val > 2.0 else PatternFill(fill_type=None))

            # Auto-fit column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or "")) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
            wb.remove(default_sheet)

        wb.save(self.export_path)
        print(f"Exported screener results to {self.export_path}")


if __name__ == "__main__":
    engine = ScreenerEngine()
    results = engine.run_all_presets()
    exporter = ScreenerExporter()
    exporter.export(results)
